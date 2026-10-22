"""Embedded lookup sources, native recalculation and saved-result verification."""
import math
import os
import shutil
import subprocess
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from .seed_reference import Workbook as SavedWorkbook, NS, RID

MAIN = NS['m']


def populate_dependencies(workbook, payload):
    """Embed the validated inputs at the addresses used by client formulas.

    Raw rows begin at 4 to match the CLIENT VLOOKUP range. The supplied seed
    starts at row 3; every validated record is copied, including its first row.
    The helper sheets are visible and editable for formula traceability.
    """
    raw = workbook.create_sheet('FGT Data Download - New')
    for address, value in {'A3': 'Lookup key', 'B3': 'StartDate', 'E3': 'LOC',
                           'N3': 'EnergyQty', 'Q3': 'BTUFctr'}.items():
        raw[address] = value
    for row, record in enumerate(payload['raw_records'], 4):
        raw[f'A{row}'] = f'=B{row}&E{row}'
        raw[f'B{row}'] = date.fromisoformat(record['date'])
        raw[f'B{row}'].number_format = 'mm/dd/yyyy'
        raw[f'E{row}'] = int(record['station_id'])
        raw[f'N{row}'] = float(Decimal(str(record['energy_qty_dt'])))
        raw[f'Q{row}'] = float(Decimal(str(record['btu_factor'])))
    total = workbook.create_sheet('TOTAL FGT OA')
    first = date.fromisoformat(payload['days'][0]['date'])
    total['A4'] = payload['source'].get('period_label') or first.strftime('%B %Y')
    gulf = workbook.create_sheet('GSNG Imbalance')
    gulf['Y10'] = 'Validated Gulfstream peaker burn (DT)'
    for index, day in enumerate(payload['days']):
        row = index + 12
        total[f'B{row}'] = date.fromisoformat(day['date'])
        total[f'B{row}'].number_format = 'mm/dd/yyyy'
        # D/E are supporting snapshots. The client's I_FPC OA D/E cells
        # themselves are inputs rather than formulas.
        total[f'D{row}'] = day.get('alert_day')
        total[f'E{row}'] = day.get('alert_direction')
        gulf[f'Y{index + 11}'] = float(Decimal(str(day['gulfstream_peaker_burn_dt'])))


def recalculate(path, engine=None):
    """Use desktop Excel on Windows, or explicitly available LibreOffice.

    No Python expression evaluator or synthetic formula caches are used.
    Set FPC_OA_RECALC_ENGINE=excel|libreoffice. Default auto selects Excel on
    Windows and LibreOffice elsewhere. Missing engines fail with instructions.
    """
    engine = engine or os.environ.get('FPC_OA_RECALC_ENGINE', 'auto')
    if engine == 'auto':
        engine = 'excel' if os.name == 'nt' else 'libreoffice'
    if engine == 'excel':
        if os.name != 'nt':
            raise RuntimeError('Excel recalculation requires Windows and desktop Excel')
        executable = shutil.which('powershell.exe')
        if not executable:
            raise RuntimeError('PowerShell is required for Excel recalculation')
        command = [executable, '-NoProfile', '-NonInteractive', '-File',
                   str(Path(__file__).with_name('recalculate_excel.ps1')), str(path.resolve())]
        result = subprocess.run(command, capture_output=True, text=True, timeout=300)
        if result.returncode:
            raise RuntimeError(f'Excel recalculation failed: {result.stderr.strip()}')
        return path, 'Microsoft Excel'
    if engine != 'libreoffice':
        raise ValueError('FPC_OA_RECALC_ENGINE must be auto, excel or libreoffice')
    executable = os.environ.get('FPC_OA_SOFFICE') or shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise RuntimeError('Formula output requires a calculation engine. On Windows use desktop '
                           'Excel; otherwise configure an approved LibreOffice executable with '
                           'FPC_OA_RECALC_ENGINE=libreoffice and FPC_OA_SOFFICE.')
    converted = path.parent / 'recalculated'
    converted.mkdir()
    profile = (path.parent / 'lo-profile').resolve().as_uri()
    command = [executable, f'-env:UserInstallation={profile}', '--headless',
               '--convert-to', 'xlsx', '--outdir', str(converted), str(path)]
    result = subprocess.run(command, capture_output=True, text=True, timeout=300)
    target = converted / path.name
    if result.returncode or not target.is_file():
        raise RuntimeError(f'LibreOffice recalculation failed: {result.stderr.strip()} '
                           f'{result.stdout.strip()}')
    return target, 'LibreOffice Calc'


def verify_formula_results(path, result, formulas, title):
    """Check every target formula's native saved value against Python results."""
    from .output import cell_updates
    expected = cell_updates(result)
    expected['A1'] = title
    actual = SavedWorkbook(path)
    differences = []
    for address in formulas:
        value = actual.get('I_FPC OA', address)
        target = expected[address]
        if isinstance(target, date):
            if isinstance(value, str):
                equal = value == target.strftime('%m/%d/%Y')
            else:
                equal = value == (target - date(1899, 12, 30)).days
        elif target is None:
            equal = value in (None, '')
        elif isinstance(target, (Decimal, int, float)):
            equal = isinstance(value, (int, float)) and math.isclose(
                float(target), value, rel_tol=1e-10, abs_tol=1e-6)
        else:
            equal = value == target
        if not equal:
            differences.append({'cell': address, 'expected': str(target), 'native': value})
    if differences:
        raise ValueError(f'Native formula results differ from Python: {differences[:10]}')
    return {'formula_results_checked': len(formulas), 'differences': [],
            'matches_python': True}


def _worksheet_locations(archive):
    book = ET.fromstring(archive.read('xl/workbook.xml'))
    rels = {r.get('Id'): r.get('Target') for r in
            ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
    import posixpath
    return {s.get('name'): (rels[s.get(RID)].lstrip('/') if rels[s.get(RID)].startswith('/')
                          else posixpath.normpath('xl/' + rels[s.get(RID)]))
            for s in book.find('m:sheets', NS)}


def preserve_formulas_with_native_caches(original, recalculated, destination):
    """Copy only genuine engine-calculated caches into our original OOXML.

    Calc may normalize formula syntax when saving. Keeping our original <f>
    nodes preserves the client's expressions (with documented internal links).
    Number formats/layout remain those generated from the client snapshot.
    """
    native = SavedWorkbook(recalculated)
    changes = {}
    count = 0
    with ZipFile(original) as source:
        for sheet, location in _worksheet_locations(source).items():
            tree = ET.fromstring(source.read(location))
            for cell in tree.findall('.//m:sheetData/m:row/m:c', NS):
                if cell.find('m:f', NS) is None:
                    continue
                value = native.get(sheet, cell.get('r'))
                if value is None:
                    # Only a cached empty string is a valid blank formula result.
                    kind = native.sheets[sheet][cell.get('r')][2]
                    if kind != 'str':
                        raise ValueError(f'Missing native cache at {sheet}!{cell.get("r")}')
                for child in list(cell):
                    if child.tag in (f'{{{MAIN}}}v', f'{{{MAIN}}}is'):
                        cell.remove(child)
                cell.attrib.pop('t', None)
                cache = ET.SubElement(cell, f'{{{MAIN}}}v')
                if isinstance(value, str) or value is None:
                    cell.set('t', 'str')
                    cache.text = value or ''
                else:
                    cache.text = repr(value)
                count += 1
            changes[location] = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
        with ZipFile(destination, 'w') as target:
            for info in source.infolist():
                target.writestr(info, changes.get(info.filename, source.read(info.filename)))
    return count


def finalize_formula_workbook(workbook, output, result, formulas, *, recalc_engine=None):
    """Publish only after native recalculation, formula and value checks pass."""
    from .client_formulas import CLIENT_FORMULAS
    from openpyxl import load_workbook
    from openpyxl.workbook.properties import CalcProperties
    workbook.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True,
                                          forceFullCalc=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='fpc-formulas-', dir=output.parent) as temp:
        root = Path(temp)
        original = root / 'original.xlsx'
        work = root / 'calculate.xlsx'
        final = root / 'verified.xlsx'
        workbook.save(original)
        shutil.copyfile(original, work)
        native_path, engine = recalculate(work, recalc_engine)
        title = workbook['TOTAL FGT OA']['A4'].value
        verification = verify_formula_results(native_path, result, formulas, title)
        cache_count = preserve_formulas_with_native_caches(original, native_path, final)
        verify_formula_results(final, result, formulas, title)
        check = load_workbook(final, read_only=True, data_only=False)
        try:
            sheet = check['I_FPC OA']
            saved = {c.coordinate: c.value for row in sheet for c in row if c.data_type == 'f'}
            if saved != formulas:
                raise ValueError('Saved formula text differs from client formula contract')
        finally:
            check.close()
        # Atomic publication without replacing an existing output file.
        os.link(final, output)
    return {'file': str(output), 'target_sheet': 'I_FPC OA', 'template_used': False,
            'vba_executed': False, 'downstream_recalculated': False,
            'formulas_written': len(formulas), 'native_formula_caches': cache_count,
            'recalculation_engine': engine, 'formula_validation': verification,
            'external_references_internalized': sum('[17]' in f for f in CLIENT_FORMULAS.values()),
            'formula_scope': 'Client I_FPC OA formulas; [17] prefix replaced with internal sheets',
            'dependency_scope': 'Validated snapshots of raw FGT records, daily dates and GSNG peakers; '
                                'not complete upstream client workflows'}
