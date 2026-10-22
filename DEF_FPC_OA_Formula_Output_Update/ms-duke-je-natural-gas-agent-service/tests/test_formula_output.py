"""Native formula tests; require installed Excel or LibreOffice."""
import copy
import os
import shutil
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from def_fpc_oa.calculations import calculate
from def_fpc_oa.client_formulas import CLIENT_FORMULAS, formula_updates
from def_fpc_oa.output import write_output
from def_fpc_oa.seed_reference import Workbook as SavedWorkbook
from def_i_imbalance.inputs import read_i_fpc_oa
from def_i_imbalance.calculations import calculate_i_imbalance
from test_calculations import example


def full_month():
    payload = example()
    daily = payload['days'][0]
    raw = payload['raw_records']
    payload['days'], payload['raw_records'] = [], []
    for index in range(31):
        day = (date(2026, 8, 1) + timedelta(days=index)).isoformat()
        payload['days'].append(dict(daily, date=day, alert_day=None))
        for record in raw:
            item = dict(record, date=day)
            if record['station_id'] == '7':
                item['energy_qty_dt'] = '0'
            payload['raw_records'].append(item)
    return payload


NATIVE_AVAILABLE = (os.name == 'nt' or shutil.which('soffice') or shutil.which('libreoffice')
                    or os.environ.get('FPC_OA_SOFFICE'))


@unittest.skipUnless(NATIVE_AVAILABLE, 'A native spreadsheet engine is required')
class NativeFormulaOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.payload = full_month()
        cls.result = calculate(cls.payload)
        cls.path = Path(cls.temp.name) / 'I_FPC_OA.xlsx'
        cls.info = write_output(cls.path, cls.result, validated_input=cls.payload)

    def test_all_client_formulas_and_no_external_links(self):
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            actual = {c.coordinate: c.value for row in workbook['I_FPC OA']
                      for c in row if c.data_type == 'f'}
            self.assertEqual(len(CLIENT_FORMULAS), 826)
            self.assertEqual(actual, formula_updates())
            for address, formula in actual.items():
                self.assertEqual(formula, CLIENT_FORMULAS[address].replace("'[17]", "'"))
            self.assertEqual(len(workbook.sheetnames), 4)
            self.assertEqual(workbook['FGT Data Download - New']['A4'].value, '=B4&E4')
        finally:
            workbook.close()
        self.assertEqual(self.info['formula_validation']['formula_results_checked'], 826)

    def test_native_caches_and_next_stage_handoff(self):
        def opener(path):
            return load_workbook(path, read_only=True, data_only=True)
        payload = read_i_fpc_oa(self.path, workbook_opener=opener)
        self.assertEqual(len(payload['records']), 31)
        before = calculate_i_imbalance(self.result)
        after = calculate_i_imbalance(payload)
        self.assertEqual(before['summary'], after['summary'])
        cached = SavedWorkbook(self.path)
        self.assertEqual(cached.get('I_FPC OA', 'F13'), -10)
        self.assertEqual(cached.get('I_FPC OA', 'L13'), 0)
        self.assertIn(cached.get('I_FPC OA', 'X13'), (None, ''))
        self.assertEqual(cached.get('I_FPC OA', 'P13'), 'NA')

    def test_changed_source_recalculates_and_preserves_first_raw_record(self):
        payload = copy.deepcopy(self.payload)
        payload['days'][0]['scheduled_volume_dt'] = '90'
        payload['days'][0]['alert_day'] = '0.05'
        payload['days'][1]['alert_day'] = ' '
        payload['raw_records'][0]['energy_qty_dt'] = '-20'
        result = calculate(payload)
        path = Path(self.temp.name) / 'changed.xlsx'
        write_output(path, result, validated_input=payload)
        actual = SavedWorkbook(path)
        self.assertEqual(actual.get('I_FPC OA', 'F13'), -20)
        self.assertEqual(actual.get('I_FPC OA', 'P13'), 4.5)
        self.assertEqual(actual.get('I_FPC OA', 'Z13'), 1023)
        self.assertEqual(actual.get('I_FPC OA', 'L47'), float(result['summary']['ending_imbalance_dt']))

    def test_invalid_engine_does_not_publish_output(self):
        path = Path(self.temp.name) / 'must-not-exist.xlsx'
        with self.assertRaisesRegex(ValueError, 'RECALC_ENGINE'):
            write_output(path, self.result, validated_input=self.payload, recalc_engine='invalid')
        self.assertFalse(path.exists())

    def test_mismatched_input_is_rejected(self):
        changed = copy.deepcopy(self.payload)
        changed['days'][0]['scheduled_volume_dt'] = '90'
        path = Path(self.temp.name) / 'mismatched.xlsx'
        with self.assertRaisesRegex(ValueError, 'does not match'):
            write_output(path, self.result, validated_input=changed)
        self.assertFalse(path.exists())


if __name__ == '__main__':
    unittest.main()
