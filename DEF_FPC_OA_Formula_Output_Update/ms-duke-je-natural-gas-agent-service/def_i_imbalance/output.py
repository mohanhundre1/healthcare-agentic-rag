"""Template-free FGT section and daily handoff; no other pipelines are invented."""

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .calculations import CANONICAL_STATIONS


# Client FGT heat headers have a different order from burn headers.
HEAT_STATIONS = (
    "Hines",
    "Vandolah",
    "Suwannee",
    "Bartow",
    "Anclote",
    "Tiger Bay",
    "DeBary",
    "Univ of FL",
)

HEAT_LABELS = (
    "Hines/Others",
    "Vandolah",
    "Suwannee",
    "Bartow CC",
    "Anclote",
    "Tiger Bay",
    "DeBary",
    "Univ of FL",
)


def write_output(output: Path, result: dict):
    output = Path(output)

    if output.exists():
        raise FileExistsError(
            f"Output already exists: {output}"
        )

    if output.suffix.lower() != ".xlsx":
        raise ValueError(
            "Template-free I_Imbalance output must use .xlsx"
        )

    wb = Workbook()

    ws = wb.active
    ws.title = "I_Imbalance"

    ws["A1"] = (
        f"{result['reporting_period']} | FGT ONLY | "
        + (
            "mapped"
            if result["complete_fgt"]
            else "requires review"
        )
    )

    ws["A4"], ws["B4"] = (
        "Source",
        "MMBTU-FGT",
    )

    ws["A5"], ws["A11"] = (
        "I_FPC OA",
        "I_FPC OA",
    )

    ws["A10"], ws["B10"] = (
        "Source",
        "HEAT RATE-FGT",
    )

    monthly = {
        row["station"]: row
        for row in result["monthly_stations"]
    }

    def put(address, value):
        """
        Write a value to an Excel cell.

        Excel uses floating-point numeric cells;
        exact Decimal values remain in JSON.
        """

        ws[address] = (
            float(value)
            if isinstance(value, Decimal)
            else value
        )

        if isinstance(
            value,
            (Decimal, int, float),
        ):
            ws[address].number_format = "#,##0.000000"

    # ---------------------------------------------------------
    # FGT monthly burn section
    # ---------------------------------------------------------

    for col, station in enumerate(
        CANONICAL_STATIONS,
        2,
    ):
        letter = get_column_letter(col)

        ws[f"{letter}5"] = station

        put(
            f"{letter}6",
            monthly[station]["MMBTU-FGT"],
        )

    # ---------------------------------------------------------
    # FGT monthly heat-rate section
    # ---------------------------------------------------------

    for col, station, label in zip(
        range(2, 10),
        HEAT_STATIONS,
        HEAT_LABELS,
    ):
        letter = get_column_letter(col)

        ws[f"{letter}11"] = label

        item = monthly[station]

        factor = item["HEAT RATE-FGT"]

        # Excel's original direct reference to an all-blank
        # factor total displays 0.
        #
        # Keep daily blanks in JSON and Daily_FGT;
        # never substitute a nonzero burn.
        if (
            factor is None
            and item["MMBTU-FGT"] == 0
            and result["complete_fgt"]
        ):
            factor = 0

        put(
            f"{letter}12",
            factor,
        )

    # ---------------------------------------------------------
    # Scope and control information
    # ---------------------------------------------------------

    ws["A15"] = (
        "Gulfstream / Osprey / Citrus / Sabal Trail "
        "inputs are not supplied by this flow."
    )

    ws["A16"] = (
        "This workbook contains the FGT section only; "
        "the full imbalance workflow is incomplete."
    )

    ws["A18"] = "Signed FGT total"

    ws["B18"] = (
        float(result["summary"]["total_fgt_mmbtu"])
        if result["summary"]["total_fgt_mmbtu"] is not None
        else None
    )

    ws["A19"] = "Exception count"

    ws["B19"] = len(
        result["exceptions"]
    )

    # ---------------------------------------------------------
    # I_Imbalance worksheet formatting
    # ---------------------------------------------------------

    ws.freeze_panes = "B5"

    ws.column_dimensions["A"].width = 24

    for col in range(2, 10):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 20

    for row in (5, 11):
        for cell in ws[row][:9]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="17365D",
            )

    # ---------------------------------------------------------
    # Daily FGT worksheet
    # ---------------------------------------------------------

    daily = wb.create_sheet(
        "Daily_FGT"
    )

    daily.append(
        [
            "Date",
            "Station",
            "Pipeline",
            "MMBTU-FGT",
            "HEAT RATE-FGT",
        ]
    )

    for record in result["records"]:
        for station in record["stations"]:
            daily.append(
                [
                    record["date"],
                    station["station"],
                    "FGT",
                ]
                + [
                    (
                        float(station[key])
                        if isinstance(
                            station[key],
                            Decimal,
                        )
                        else station[key]
                    )
                    for key in (
                        "MMBTU-FGT",
                        "HEAT RATE-FGT",
                    )
                ]
            )

    # ---------------------------------------------------------
    # Daily FGT worksheet formatting
    # ---------------------------------------------------------

    daily.freeze_panes = "A2"

    daily.auto_filter.ref = daily.dimensions

    for col in (
        "A",
        "B",
        "C",
        "D",
        "E",
    ):
        daily.column_dimensions[col].width = 22

    # ---------------------------------------------------------
    # Save generated workbook
    # ---------------------------------------------------------

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb.save(
        output
    )

    wb.close()

    # ---------------------------------------------------------
    # Return output metadata
    # ---------------------------------------------------------

    return {
        "file": output.name,
        "target_sheet": "I_Imbalance",
        "template_used": False,
        "scope": "FGT_ONLY",
        "complete_workflow": False,
    }