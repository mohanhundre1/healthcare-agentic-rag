"""Regression cases absent from the original all-nonzero synthetic fixture."""

import copy
import json
import tempfile
import unittest

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from test_i_imbalance_calculations import example_i_fpc_oa

from def_i_imbalance.calculations import calculate_i_imbalance
from def_i_imbalance.inputs import (
    _find_calculated_payload,
    read_i_fpc_oa,
)
from def_i_imbalance.output import write_output
from def_i_imbalance.workflow import process_i_imbalance


class IImbalanceRegressions(unittest.TestCase):

    def zero_burn(self):
        p = example_i_fpc_oa()

        p["records"][0]["burns_dt"]["tiger bay"] = "0"
        p["records"][0]["btu_factors"]["tiger bay"] = None

        return p

    def test_zero_burn_blank_factor_is_valid(self):
        r = calculate_i_imbalance(
            self.zero_burn()
        )

        self.assertTrue(
            r["complete_fgt"]
        )

        self.assertEqual(
            r["exceptions"],
            [],
        )

        tiger = r["records"][0]["stations"][6]

        self.assertEqual(
            tiger["MMBTU-FGT"],
            0,
        )

        self.assertIsNone(
            tiger["HEAT RATE-FGT"]
        )

    def test_nonzero_burn_blank_factor_still_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["btu_factors"]["hines"] = None

        with self.assertRaisesRegex(
            ValueError,
            "MISSING_FGT_HEAT_RATE",
        ):
            calculate_i_imbalance(p)

    def test_blank_zero_factor_does_not_hide_invalid_number(self):
        p = self.zero_burn()

        p["records"][0]["btu_factors"]["tiger bay"] = "NaN"

        with self.assertRaisesRegex(
            ValueError,
            "finite",
        ):
            calculate_i_imbalance(p)

    def test_boolean_burn_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["burns_dt"]["hines"] = True

        with self.assertRaisesRegex(
            ValueError,
            "Boolean",
        ):
            calculate_i_imbalance(p)

    def test_monthly_averages_exclude_zero_burn_blank_days(self):
        p = self.zero_burn()

        second = copy.deepcopy(
            p["records"][0]
        )

        second["date"] = "2026-08-02"
        second["burns_dt"]["tiger bay"] = "-5"
        second["btu_factors"]["tiger bay"] = "1022"

        p["records"].append(second)

        r = calculate_i_imbalance(p)

        tiger = r["monthly_stations"][6]

        self.assertEqual(
            tiger["MMBTU-FGT"],
            -5,
        )

        self.assertEqual(
            tiger["HEAT RATE-FGT"],
            1022,
        )

        self.assertEqual(
            r["summary"]["target_station_records"],
            16,
        )

    def test_invalid_calendar_date_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["date"] = "2026-08-99"

        with self.assertRaisesRegex(
            ValueError,
            "Invalid",
        ):
            calculate_i_imbalance(p)

    def test_gap_rejected(self):
        p = example_i_fpc_oa()

        second = copy.deepcopy(
            p["records"][0]
        )

        second["date"] = "2026-08-03"

        p["records"].append(second)

        with self.assertRaisesRegex(
            ValueError,
            "contiguous",
        ):
            calculate_i_imbalance(p)

    def test_cross_month_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["date"] = "2026-08-31"

        second = copy.deepcopy(
            p["records"][0]
        )

        second["date"] = "2026-09-01"

        p["records"].append(second)

        with self.assertRaisesRegex(
            ValueError,
            "one reporting month",
        ):
            calculate_i_imbalance(p)

    def test_datetime_normalized_to_day(self):
        p = example_i_fpc_oa()

        p["records"][0]["date"] = datetime(
            2026,
            8,
            1,
            12,
        )

        result = calculate_i_imbalance(p)

        self.assertEqual(
            result["records"][0]["date"],
            "2026-08-01",
        )

    def test_duplicate_station_alias_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["burns_dt"]["gainesville"] = "-10"

        with self.assertRaisesRegex(
            ValueError,
            "Duplicate source mapping",
        ):
            calculate_i_imbalance(p)

    def test_unknown_station_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["burns_dt"]["unknown"] = "-10"

        with self.assertRaisesRegex(
            ValueError,
            "UNMAPPED_FGT_STATION",
        ):
            calculate_i_imbalance(p)

    def test_non_strict_missing_burn_keeps_total_unavailable(self):
        p = self.zero_burn()

        del p["records"][0]["burns_dt"]["hines"]

        r = calculate_i_imbalance(
            p,
            strict=False,
        )

        self.assertFalse(
            r["complete_fgt"]
        )

        self.assertIsNone(
            r["summary"]["total_fgt_mmbtu"]
        )

        self.assertEqual(
            r["exceptions"][0]["date"],
            "2026-08-01",
        )

    def test_bad_monthly_control_rejected(self):
        p = example_i_fpc_oa()

        p["summary"]["total_plant_burns_dt"] = "-999"

        with self.assertRaisesRegex(
            ValueError,
            "monthly burns",
        ):
            calculate_i_imbalance(p)

    def test_bad_daily_control_rejected(self):
        p = example_i_fpc_oa()

        p["records"][0]["total_plant_burns_dt"] = "-999"

        with self.assertRaisesRegex(
            ValueError,
            "daily burns",
        ):
            calculate_i_imbalance(p)

    def test_upstream_exceptions_not_discarded(self):
        p = example_i_fpc_oa()

        p["exceptions"] = [
            {
                "kind": "missing_lookup"
            }
        ]

        with self.assertRaisesRegex(
            ValueError,
            "upstream exceptions",
        ):
            calculate_i_imbalance(p)

        r = calculate_i_imbalance(
            p,
            strict=False,
        )

        self.assertFalse(
            r["complete_fgt"]
        )

    def test_upstream_rollforward_failure_rejected(self):
        p = example_i_fpc_oa()

        p["control"] = {
            "rollforward_balanced": False
        }

        with self.assertRaisesRegex(
            ValueError,
            "rollforward",
        ):
            calculate_i_imbalance(p)

    def test_schema_rejected(self):
        p = example_i_fpc_oa()

        p["schema_version"] = "2"

        with self.assertRaisesRegex(
            ValueError,
            "schema_version",
        ):
            calculate_i_imbalance(p)

    def test_ending_imbalance_has_consistent_numeric_type(self):
        p = self.zero_burn()

        text_result = calculate_i_imbalance(p)

        p["summary"]["ending_imbalance_dt"] = 100

        numeric_result = calculate_i_imbalance(p)

        self.assertEqual(
            text_result["summary"],
            numeric_result["summary"],
        )

        self.assertIsInstance(
            text_result["summary"]["ending_imbalance_dt"],
            Decimal,
        )

    def test_nested_a2a_result_list(self):
        p = self.zero_burn()

        result = _find_calculated_payload(
            [
                {
                    "result": {
                        "data": p
                    }
                }
            ]
        )

        self.assertIs(
            result,
            p,
        )

    def test_multiple_results_are_ambiguous(self):
        with self.assertRaisesRegex(
            ValueError,
            "exactly one",
        ):
            _find_calculated_payload(
                [
                    self.zero_burn(),
                    self.zero_burn(),
                ]
            )

    def test_bom_json_input(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.json"

            path.write_text(
                json.dumps(
                    self.zero_burn()
                ),
                encoding="utf-8-sig",
            )

            result = read_i_fpc_oa(path)

            self.assertEqual(
                result["records"][0]["date"],
                "2026-08-01",
            )

    def test_non_strict_bad_heat_cannot_produce_valid_monthly_average(
        self,
    ):
        p = self.zero_burn()

        second = copy.deepcopy(
            p["records"][0]
        )

        second["date"] = "2026-08-02"
        second["btu_factors"]["hines"] = None

        p["records"].append(second)

        r = calculate_i_imbalance(
            p,
            strict=False,
        )

        self.assertIsNone(
            r["monthly_stations"][3]["HEAT RATE-FGT"]
        )

    def test_workflow_returns_real_xlsx_and_explicit_fgt_scope(self):
        import base64

        from io import BytesIO
        from openpyxl import load_workbook

        r = process_i_imbalance(
            {
                "i_fpc_oa": self.zero_burn()
            }
        )

        workbook_bytes = base64.b64decode(
            r["output"]["workbook_base64"]
        )

        wb = load_workbook(
            BytesIO(workbook_bytes),
            data_only=True,
        )

        try:
            self.assertEqual(
                wb["I_Imbalance"]["B6"].value,
                -10,
            )

            self.assertEqual(
                wb["I_Imbalance"]["B11"].value,
                "Hines/Others",
            )

            self.assertEqual(
                wb["I_Imbalance"]["G12"].value,
                0,
            )

            self.assertEqual(
                wb["Daily_FGT"].max_row,
                9,
            )

            self.assertFalse(
                r["complete_workflow"]
            )

            self.assertEqual(
                r["scope"],
                "FGT_ONLY",
            )

        finally:
            wb.close()

    def test_output_overwrite_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "I_Imbalance.xlsx"

            path.write_bytes(
                b"original"
            )

            with self.assertRaises(
                FileExistsError
            ):
                write_output(
                    path,
                    calculate_i_imbalance(
                        self.zero_burn()
                    ),
                )

            self.assertEqual(
                path.read_bytes(),
                b"original",
            )


if __name__ == "__main__":
    unittest.main()