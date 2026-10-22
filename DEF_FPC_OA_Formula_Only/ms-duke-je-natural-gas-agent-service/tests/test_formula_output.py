"""Formula-only output and calculated-JSON handoff checks."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from def_fpc_oa.calculations import calculate, jsonable
from def_fpc_oa.client_formulas import formula_updates
from def_fpc_oa.output import write_output
from def_i_imbalance.inputs import read_i_fpc_oa
from def_i_imbalance.calculations import calculate_i_imbalance
from test_calculations import example
from datetime import date, timedelta

def full_month():
    payload = example()
    daily = payload['days'][0]
    raw = payload['raw_records']
    payload['days'], payload['raw_records'] = [], []
    for index in range(31):
        day = (date(2026, 8, 1) + timedelta(days=index)).isoformat()
        payload['days'].append(dict(daily, date=day, alert_day=None))
        for record in raw:
            item = dict(record, date=day)
            if record['station_id'] == '7':
                item['energy_qty_dt'] = '0'
            payload['raw_records'].append(item)
    return payload


class FormulaOnlyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.payload = full_month()
        cls.result = calculate(cls.payload)
        cls.path = Path(cls.temp.name) / 'I_FPC_OA.xlsx'
        with patch('subprocess.run', side_effect=AssertionError('No external engine allowed')):
            cls.info = write_output(cls.path, cls.result, validated_input=cls.payload)

    def test_formula_text_is_retained_without_recalculation(self):
        workbook = load_workbook(self.path, read_only=True, data_only=False)
        try:
            actual = {c.coordinate: c.value for row in workbook['I_FPC OA']
                      for c in row if c.data_type == 'f'}
            self.assertEqual(actual, formula_updates())
            self.assertEqual(len(actual), 826)
            self.assertEqual(workbook.calculation.calcMode, 'manual')
            self.assertFalse(workbook.calculation.fullCalcOnLoad)
        finally:
            workbook.close()
        self.assertFalse(self.info['recalculation_performed'])

    def test_no_synthetic_formula_results(self):
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            self.assertIsNone(workbook['I_FPC OA']['N13'].value)
            self.assertIsNone(workbook['I_FPC OA']['F13'].value)
            self.assertEqual(workbook['I_FPC OA']['C13'].value, 80)
        finally:
            workbook.close()
        def opener(path):
            return load_workbook(path, read_only=True, data_only=True)
        with self.assertRaisesRegex(ValueError, 'Missing saved formula value'):
            read_i_fpc_oa(self.path, workbook_opener=opener)

    def test_json_handoff_preserves_next_stage_results(self):
        path = Path(self.temp.name) / 'calculated_I_FPC_OA.json'
        path.write_text(json.dumps(jsonable(self.result)), encoding='utf-8')
        actual = calculate_i_imbalance(read_i_fpc_oa(path))
        expected = calculate_i_imbalance(self.result)
        self.assertEqual(actual['summary'], expected['summary'])

    def test_mismatched_inputs_do_not_publish_output(self):
        payload = copy.deepcopy(self.payload)
        payload['days'][0]['scheduled_volume_dt'] = '90'
        path = Path(self.temp.name) / 'bad.xlsx'
        with self.assertRaisesRegex(ValueError, 'does not match'):
            write_output(path, self.result, validated_input=payload)
        self.assertFalse(path.exists())
